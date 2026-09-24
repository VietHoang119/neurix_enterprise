"""Turn an uploaded file into plain text for the ingest flow.

PDF and images go straight to Gemini (inline bytes, no local parsing);
docx/xlsx are read with python-docx/openpyxl; txt/md are decoded as UTF-8.
Anything else is rejected with a user-facing message.
"""

import io

from langchain_core.messages import HumanMessage

from core import llm

_TEXT_TYPES = {"text", "markdown", "md", "txt", "plain"}
_GEMINI_PROMPT = (
    "Extract the full content of this file as plain text so it can be saved as a note. "
    "Transcribe all readable text faithfully, keeping its structure (headings, lists, tables as rows). "
    "If it is an image, also describe what it shows. Return only the extracted content, no commentary."
)


class UnsupportedFileError(ValueError):
    pass


def file_kind(file_info: dict) -> str:
    """Classify a Slack files.info record: 'pdf', 'image', 'docx', 'xlsx' or 'text'."""
    filetype = (file_info.get("filetype") or "").lower()
    mimetype = (file_info.get("mimetype") or "").lower()
    name = (file_info.get("name") or "").lower()
    if filetype == "pdf" or mimetype == "application/pdf":
        return "pdf"
    if mimetype.startswith("image/"):
        return "image"
    if filetype == "docx" or name.endswith(".docx"):
        return "docx"
    if filetype == "xlsx" or name.endswith(".xlsx"):
        return "xlsx"
    if filetype in _TEXT_TYPES or mimetype.startswith("text/") or name.endswith((".txt", ".md")):
        return "text"
    raise UnsupportedFileError(
        f"Unsupported file type ({filetype or mimetype or 'unknown'}). "
        "Supported: PDF, images, .docx, .xlsx, .txt, .md."
    )


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                lines.append("\t".join(cells))
    return "\n".join(lines)


async def _gemini_text(data: bytes, mime_type: str) -> str:
    message = HumanMessage(
        content=[
            {"type": "text", "text": _GEMINI_PROMPT},
            {"type": "media", "mime_type": mime_type, "data": data},
        ]
    )
    response = await llm.get_chat_model().ainvoke([message])
    content = response.content
    if isinstance(content, list):
        content = "".join(part.get("text", "") if isinstance(part, dict) else str(part) for part in content)
    return content or ""


async def extract_text(file_info: dict, data: bytes) -> str:
    kind = file_kind(file_info)
    if kind == "pdf":
        text = await _gemini_text(data, "application/pdf")
    elif kind == "image":
        text = await _gemini_text(data, file_info.get("mimetype") or "image/png")
    elif kind == "docx":
        text = _docx_text(data)
    elif kind == "xlsx":
        text = _xlsx_text(data)
    else:
        text = data.decode("utf-8", errors="replace")
    text = text.strip()
    if not text:
        raise UnsupportedFileError("No readable content found in that file.")
    return text
